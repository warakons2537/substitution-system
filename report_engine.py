from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path

import pandas as pd

THAI_MONTH_NAMES = [
    "", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม",
]
THAI_DAYS = ["จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์", "เสาร์", "อาทิตย์"]


def thai_date_long(d: date) -> str:
    return f"วัน{THAI_DAYS[d.weekday()]} ที่ {d.day} {THAI_MONTH_NAMES[d.month]} พ.ศ. {d.year + 543}"


def _checkbox(selected: bool) -> str:
    return "☒" if selected else "☐"


def make_substitution_docx(
    daydf: pd.DataFrame,
    selected_date: date,
    absent_teacher: str,
    settings: dict[str, str],
) -> bytes:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.4)
    sec.bottom_margin = Cm(1.4)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "TH Sarabun New"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    normal.font.size = Pt(16)

    def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=15):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = align
        r = p.add_run(str(text or ""))
        r.bold = bold
        r.font.name = "TH Sarabun New"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
        r.font.size = Pt(size)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("บันทึกการสอนแทน")
    r.bold = True; r.font.name = "TH Sarabun New"; r.font.size = Pt(20)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(settings.get("school_name", ""))
    r.bold = True; r.font.name = "TH Sarabun New"; r.font.size = Pt(17)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(thai_date_long(selected_date))
    r.font.name = "TH Sarabun New"; r.font.size = Pt(16)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")

    p = doc.add_paragraph()
    r = p.add_run(f"ชื่อ ครูประจำวิชา  {absent_teacher}")
    r.font.name = "TH Sarabun New"; r.font.size = Pt(16)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")

    leave_type = "ไม่ระบุ"
    if not daydf.empty and "leave_type" in daydf.columns:
        vals = [str(x) for x in daydf["leave_type"].dropna().tolist() if str(x).strip()]
        if vals:
            leave_type = vals[0]
    labels = ["ลาป่วย", "ลากิจ", "ไปราชการ", "ลาคลอด"]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, label in enumerate(labels):
        r = p.add_run(f"{_checkbox(leave_type == label)} {label}")
        r.font.name = "TH Sarabun New"; r.font.size = Pt(15)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
        if i < len(labels)-1:
            p.add_run("      ")

    headers = ["ชั่วโมงที่", "วิชา", "ชั้น", "ผู้สอนแทน", "รับทราบ", "หมายเหตุ"]
    periods = int(settings.get("periods_per_day", "6") or 6)
    periods = max(1, min(periods, 20))
    table = doc.add_table(rows=1 + periods, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [Cm(1.7), Cm(4.2), Cm(2.1), Cm(3.8), Cm(2.5), Cm(3.2)]
    for j, h in enumerate(headers):
        set_cell_text(table.cell(0,j), h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        table.cell(0,j).width = widths[j]

    lookup = {}
    if not daydf.empty:
        for _, row in daydf.iterrows():
            try:
                lookup[int(row.get("period", 0))] = row
            except Exception:
                pass
    for period in range(1, periods+1):
        row = lookup.get(period)
        vals = [
            period,
            "" if row is None else row.get("subject", ""),
            "" if row is None else row.get("class_name", ""),
            "" if row is None else row.get("substitute_teacher", ""),
            "" if row is None else row.get("acknowledged", ""),
            "" if row is None else row.get("note", ""),
        ]
        for j, val in enumerate(vals):
            set_cell_text(table.cell(period,j), val, align=WD_ALIGN_PARAGRAPH.CENTER if j in [0,2,4] else WD_ALIGN_PARAGRAPH.LEFT)
            table.cell(period,j).width = widths[j]
        trPr = table.rows[period]._tr.get_or_add_trPr()
        trHeight = OxmlElement('w:trHeight')
        trHeight.set(qn('w:val'), '620')
        trHeight.set(qn('w:hRule'), 'atLeast')
        trPr.append(trHeight)

    doc.add_paragraph("")
    sig = doc.add_table(rows=1, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig.autofit = False
    left = sig.cell(0,0); right = sig.cell(0,1)
    set_cell_text(left, f"ลงชื่อ ............................................................\n({absent_teacher})\nครูประจำวิชา", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(right, f"ลงชื่อ ............................................................\n({settings.get('academic_head_name','')})\nหัวหน้าบริหารงานวิชาการ", align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"ลงชื่อ ............................................................\n({settings.get('director_name','')})\nผู้อำนวยการสถานศึกษา")
    r.font.name = "TH Sarabun New"; r.font.size = Pt(15)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")

    out = io.BytesIO(); doc.save(out); return out.getvalue()


def _find_soffice() -> str | None:
    candidates = [
        shutil.which("soffice"), shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    return next((c for c in candidates if c and os.path.exists(c)), None)


def docx_bytes_to_pdf(docx_bytes: bytes) -> bytes:
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError("เครื่องนี้ยังไม่พบ LibreOffice จึงสร้าง PDF โดยตรงไม่ได้ แต่สามารถดาวน์โหลด DOCX แล้วพิมพ์จาก Word ได้")
    with tempfile.TemporaryDirectory(prefix="subreport_") as td:
        td = Path(td)
        src = td / "report.docx"; src.write_bytes(docx_bytes)
        cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(td), str(src)]
        p = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=90)
        pdf = td / "report.pdf"
        if not pdf.exists():
            raise RuntimeError("สร้าง PDF ไม่สำเร็จ: " + p.stdout[-500:])
        return pdf.read_bytes()


def make_excel(df: pd.DataFrame, title: str = "รายงานการสอนแทน") -> bytes:
    out = io.BytesIO()
    display_cols = [
        "date_iso","day_name","thai_year","absent_teacher","leave_type","period",
        "subject","class_name","substitute_teacher","acknowledged","note","source_file",
    ]
    work = df.copy()
    for c in display_cols:
        if c not in work.columns:
            work[c] = ""
    teacher_summary = (work.groupby("substitute_teacher", dropna=False).size().reset_index(name="จำนวนคาบ")
                       .sort_values("จำนวนคาบ", ascending=False)) if not work.empty else pd.DataFrame(columns=["substitute_teacher","จำนวนคาบ"])
    absent_summary = (work.groupby("absent_teacher", dropna=False).size().reset_index(name="จำนวนคาบ")
                      .sort_values("จำนวนคาบ", ascending=False)) if not work.empty else pd.DataFrame(columns=["absent_teacher","จำนวนคาบ"])
    subject_summary = (work.groupby("subject", dropna=False).size().reset_index(name="จำนวนคาบ")
                       .sort_values("จำนวนคาบ", ascending=False)) if not work.empty else pd.DataFrame(columns=["subject","จำนวนคาบ"])
    daily_summary = (work.groupby(["date_iso","absent_teacher"], dropna=False).size().reset_index(name="จำนวนคาบ")
                     .sort_values("date_iso", ascending=False)) if not work.empty else pd.DataFrame(columns=["date_iso","absent_teacher","จำนวนคาบ"])

    with pd.ExcelWriter(out, engine="xlsxwriter") as writer:
        work[display_cols].to_excel(writer, sheet_name="บันทึกข้อมูล", index=False)
        teacher_summary.to_excel(writer, sheet_name="สรุปผู้สอนแทน", index=False)
        absent_summary.to_excel(writer, sheet_name="สรุปคนลา", index=False)
        subject_summary.to_excel(writer, sheet_name="สรุปวิชา", index=False)
        daily_summary.to_excel(writer, sheet_name="สรุปรายวัน", index=False)
        wb = writer.book
        header_fmt = wb.add_format({"bold": True, "bg_color": "#2F75B5", "font_color": "#FFFFFF", "align": "center", "border": 1})
        title_fmt = wb.add_format({"bold": True, "font_size": 16, "bg_color": "#17365D", "font_color": "#FFFFFF", "align": "center"})
        for sname, ws in writer.sheets.items():
            ws.freeze_panes(1, 0)
            ws.set_row(0, 24, header_fmt)
            ws.set_column(0, 0, 16)
            ws.set_column(1, 3, 20)
            ws.set_column(4, 20, 18)
            rows, cols = (len(work), len(display_cols)) if sname == "บันทึกข้อมูล" else (max(1, ws.dim_rowmax), max(1, ws.dim_colmax + 1))
            if cols > 0:
                ws.autofilter(0, 0, max(1, rows), max(0, cols - 1))
    return out.getvalue()
